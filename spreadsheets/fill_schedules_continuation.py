# Public-repository configuration layer (replaces machine-specific paths).
from pathlib import Path as _ConfigPath
import sys as _config_sys
_PROJECT_ROOT = _ConfigPath(__file__).resolve().parents[1]
if str(_PROJECT_ROOT) not in _config_sys.path:
    _config_sys.path.insert(0, str(_PROJECT_ROOT))
from common.config import (  # noqa: E402
    WORKSPACE_PREFIX,
    ISSUE_OUTPUT_PREFIX,
    PHOTO_ARCHIVE_PREFIX,
    PHOTO_PROCESSING_PREFIX,
    CHROMEDRIVER_PATH,
    PUBLICATION_YEAR,
)

import pandas as pd
import openpyxl
import sys, os, re
import shutil
from itertools import groupby

def sort_file(input_file, output_file):
    with open(input_file) as f:
        with open(output_file, "w") as o:
            o.write("\n".join(sorted(f.read().splitlines())))

def replacer(initial_string, ch, replacing_character, occurrence):

    # breaking a string into
    # it's every single character
    lst1 = list(initial_string)
    lst2 = list(ch)

    # Loop to find the occurrence
    # of the character in the string
    # and replace it with the given
    # replacing_character
    for j in lst2:
        sub_string = j
        checklist = [i for i in range(0, len(initial_string))
              if initial_string[i:].startswith(sub_string)]
        if len(checklist)>= occurrence:
            lst1[checklist[occurrence-1]] = replacing_character

    return ''.join(lst1)

fot_tab = []
fot_spa = []
fot_pow = []


fp = open(WORKSPACE_PREFIX + "myfile3.txt", 'r', encoding="utf-8")
for l_no, line in enumerate(fp):
    if 'fotka' in line:
        fot_tab.append(line[6:])

ff = open(WORKSPACE_PREFIX + "fotki3.txt",mode='w', encoding="utf-8")
for elem in fot_tab:
    ff.write(elem)

ff.close()
fp.close()
"""
fot_tab.sort()
fot_tab2 = []
for elem in fot_tab:
    replacer(elem,"_","__",4)
    fot_tab2.append(elem)

fot_tab3 = [list(w) for z, w in groupby(fot_tab2, lambda a: a.split('__')[0])]

tab = []
# usuwanie nawiasow i przecinków
for i in fot_tab3:
    i = str(i)[1:-1]
    i = re.sub(",", "", i)
    i = re.sub("__", "_", i)
    i = re.sub("\'", "", i)
    tab.append(i)

ff = open(WORKSPACE_PREFIX + "fotki3.txt",mode='w', encoding="utf-8")
for elem in tab:
    print(str(elem))
    ff.write(str(elem))
ff.close()


with open(WORKSPACE_PREFIX + 'fotki3.txt', 'w', encoding="utf-8") as fp:
    for item in tab:
        # write each item on a new line
        fp.write("%s" % item)


ff = open(WORKSPACE_PREFIX + "fotki.txt",mode='w', encoding="utf-8")
for elem in fot_tab:
    ff.write(elem)
    if elem != None and "PowA" in elem:
        #print("Powiat: ",elem)
        sep = "IJ_"
        tekst = elem.split(sep, 1)[1]
        sep = "_"
        nr = tekst.split(sep, 1)[0]
        fot_pow.append(int(nr))
        fot_pow.append(elem)
    elif elem != None and "SpA" in elem:
        #print("Sport: ",elem)
        sep = "IJ_"
        tekst = elem.split(sep, 1)[1]
        sep = "_"
        nr = tekst.split(sep, 1)[0]
        fot_spa.append(int(nr))
        fot_spa.append(elem)
ff.close()
"""
fot_tab=[]
fot_tab2=[]
fot_tab3=[]
fot_tab4=[]
ff = open(WORKSPACE_PREFIX + "fotki3.txt",mode='r', encoding="utf-8")
fot_tab = ff.readlines()
ff.close()

fot_tab.sort()

for i in fot_tab:
    print(i)
    #i = re.sub("", "", i)
    #i = str(i)[1:-1]
    #i = str(i)[:-1]
    i = re.sub(",", "", i)
    #i = re.sub("__", "_", i)
    i = re.sub("\'", "", i)
    fot_tab3.append(i)

for el in fot_tab3:
    el = replacer(el,"_","__",4)
    fot_tab4.append(el)

fot_tab5 = [list(w) for z, w in groupby(fot_tab4, lambda a: a.split('__')[0])]

tab = []
# usuwanie nawiasow i przecinków
for i in fot_tab5:
    print(i)
    #i = re.sub("", "", i)
    #i = str(i)[1:-1]
    i = str(i)[1:-1]
    i = re.sub(",", "", i)
    i = re.sub("__", "_", i)
    i = re.sub("\'", "", i)
    tab.append(i)

for el in tab:
    print(el)


sort_file(WORKSPACE_PREFIX + "fotki3.txt", WORKSPACE_PREFIX + "fotki4.txt")
ff.close()

#f = open(WORKSPACE_PREFIX + "myfile.txt",mode='r', encoding="utf-8")
f_spa = open(WORKSPACE_PREFIX + "spa.txt",mode='r', encoding="utf-8")
f_pow = open(WORKSPACE_PREFIX + "pow.txt",mode='r', encoding="utf-8")
f_nk = open(WORKSPACE_PREFIX + "nk.txt",mode='r', encoding="utf-8")
f_kc = open(WORKSPACE_PREFIX + "kc.txt",mode='r', encoding="utf-8")
f_szb = open(WORKSPACE_PREFIX + "szb.txt",mode='r', encoding="utf-8")
f_mr = open(WORKSPACE_PREFIX + "mr.txt",mode='r', encoding="utf-8")
f_sd = open(WORKSPACE_PREFIX + "sd.txt",mode='r', encoding="utf-8")
f_rol = open(WORKSPACE_PREFIX + "rol.txt",mode='r', encoding="utf-8")
f_pol = open(WORKSPACE_PREFIX + "pol.txt",mode='r', encoding="utf-8")
f_fotki = open(WORKSPACE_PREFIX + "fotki3.txt",mode='r', encoding="utf-8")

Lines_spa = f_spa.readlines()
Lines_pow = f_pow.readlines()
Lines_nk = f_nk.readlines()
Lines_kc = f_kc.readlines()
Lines_szb = f_szb.readlines()
Lines_mr = f_mr.readlines()
Lines_sd = f_sd.readlines()
Lines_rol = f_rol.readlines()
Lines_pol = f_pol.readlines()
Lines_fotki = f_fotki.readlines()

f_spa.close()
f_pow.close()
f_nk.close()
f_szb.close()
f_mr.close()
f_kc.close()
f_sd.close()
f_rol.close()
f_pol.close()
f_fotki.close()

wp_rol = str(sys.argv[9])
wp_spa = str(sys.argv[2])
wp_powiat = str(sys.argv[3])
wp_naklo = str(sys.argv[4])
wp_szubin = str(sys.argv[5])
wp_kcynia = str(sys.argv[6])
wp_mrocza = str(sys.argv[7])
wp_sadki = str(sys.argv[8])
wp_pol = str(sys.argv[10])


#nr_gazety = "03"
nr_gazety = str(sys.argv[1])
l_art_spa = int(len(Lines_spa)/4)
l_art_pow = int(len(Lines_pow)/4)
l_art_nk = int(len(Lines_nk)/4)
l_art_kc = int(len(Lines_kc)/4)
l_art_szb = int(len(Lines_szb)/4)
l_art_mr = int(len(Lines_mr)/4)
l_art_sd = int(len(Lines_sd)/4)
l_art_rol = int(len(Lines_rol)/4)
l_art_pol = int(len(Lines_pol)/4)
# SPORT
fn = ISSUE_OUTPUT_PREFIX + ""+nr_gazety+"/"+nr_gazety+"_SpA_IJ_01_rozpiska.xlsx"
source_file = openpyxl.load_workbook(fn)
sheet = source_file["Rozpiska"]
k=0
zdj = []
brak = False
for g in tab:
    if "SpA" in g:
        zdj.append(g[:-2])
        brak = True
    k = k + 1
i=0
while not i==l_art_spa:
    sheet.cell(row=5+i+int(wp_spa), column=5).value = Lines_spa[i]
    sheet.cell(row=5+i+int(wp_spa), column=3).value = Lines_spa[i+l_art_spa*1]
    sheet.cell(row=5+i+int(wp_spa), column=10).value = int(Lines_spa[i+l_art_spa*2])
    sheet.cell(row=5+i+int(wp_spa), column=6).value = Lines_spa[i+l_art_spa*3]
    if brak:
        sheet.cell(row=5+i+int(wp_spa), column=7).value = zdj[i]
    i = i + 1
sheet.cell(row=1, column=3).value = nr_gazety+"/2026 wersja_A"
source_file.save(fn)
#print("1x1 : ", str(sheet.cell(row=1, column=3).value))
source_file.close()
# POWIAT
fn = ISSUE_OUTPUT_PREFIX + ""+nr_gazety+"/"+nr_gazety+"_PowA_IJ_01_rozpiska.xlsx"
source_file = openpyxl.load_workbook(fn)
sheet = source_file["Rozpiska"]
k=0
zdj = []
brak = False
for g in tab:
    if "PowA" in g:
        zdj.append(g[:-2])
        brak = True
    k = k + 1
i=0
while not i==l_art_pow:
    sheet.cell(row=5+i+int(wp_powiat), column=5).value = Lines_pow[i]
    sheet.cell(row=5+i+int(wp_powiat), column=3).value = Lines_pow[i+l_art_pow*1]
    sheet.cell(row=5+i+int(wp_powiat), column=10).value = int(Lines_pow[i+l_art_pow*2])
    sheet.cell(row=5+i+int(wp_powiat), column=6).value = Lines_pow[i+l_art_pow*3]
    if brak:
        sheet.cell(row=5+i+int(wp_powiat), column=7).value = zdj[i]
    i = i + 1
sheet.cell(row=1, column=3).value = nr_gazety+"/2026 wersja_A"
source_file.save(fn)
#print("1x1 : ", str(sheet.cell(row=1, column=3).value))
source_file.close()
# NAKLO
fn = ISSUE_OUTPUT_PREFIX + ""+nr_gazety+"/"+nr_gazety+"_Nk_IJ_01_rozpiska.xlsx"
source_file = openpyxl.load_workbook(fn)
sheet = source_file["Rozpiska"]
k=0
zdj = []
brak = False
for g in tab:
    if "Nk" in g:
        zdj.append(g[:-2])
        brak = True
    k = k + 1
i=0
while not i==l_art_nk:
    sheet.cell(row=5+i+int(wp_naklo), column=5).value = Lines_nk[i]
    sheet.cell(row=5+i+int(wp_naklo), column=3).value = Lines_nk[i+l_art_nk*1]
    sheet.cell(row=5+i+int(wp_naklo), column=10).value = int(Lines_nk[i+l_art_nk*2])
    sheet.cell(row=5+i+int(wp_naklo), column=6).value = Lines_nk[i+l_art_nk*3]
    if brak:
        sheet.cell(row=5+i+int(wp_naklo), column=7).value = zdj[i]
    i = i + 1
sheet.cell(row=1, column=3).value = nr_gazety+"/2026 wersja_A"
source_file.save(fn)
#print("1x1 : ", str(sheet.cell(row=1, column=3).value))
source_file.close()
# SZUBIN
fn = ISSUE_OUTPUT_PREFIX + ""+nr_gazety+"/"+nr_gazety+"_Szb_IJ_01_rozpiska.xlsx"
source_file = openpyxl.load_workbook(fn)
sheet = source_file["Rozpiska"]
k=0
zdj = []
brak = False
for g in tab:
    if "Szb" in g:
        zdj.append(g[:-2])
        brak = True
    k = k + 1
i=0
while not i==l_art_szb:
    sheet.cell(row=5+i+int(wp_szubin), column=5).value = Lines_szb[i]
    sheet.cell(row=5+i+int(wp_szubin), column=3).value = Lines_szb[i+l_art_szb*1]
    sheet.cell(row=5+i+int(wp_szubin), column=10).value = int(Lines_szb[i+l_art_szb*2])
    sheet.cell(row=5+i+int(wp_szubin), column=6).value = Lines_szb[i+l_art_szb*3]
    if brak:
        sheet.cell(row=5+i+int(wp_szubin), column=7).value = zdj[i]
    i = i + 1
sheet.cell(row=1, column=3).value = nr_gazety+"/2026 wersja_A"
source_file.save(fn)
#print("1x1 : ", str(sheet.cell(row=1, column=3).value))
source_file.close()
# MROCZA
fn = ISSUE_OUTPUT_PREFIX + ""+nr_gazety+"/"+nr_gazety+"_Mr_IJ_01_rozpiska.xlsx"
source_file = openpyxl.load_workbook(fn)
sheet = source_file["Rozpiska"]
k=0
zdj = []
brak = False
for g in tab:
    if "Mr" in g:
        zdj.append(g[:-2])
        brak = True
    k = k + 1
i=0
while not i==l_art_mr:
    sheet.cell(row=5+i+int(wp_mrocza), column=5).value = Lines_mr[i]
    sheet.cell(row=5+i+int(wp_mrocza), column=3).value = Lines_mr[i+l_art_mr*1]
    sheet.cell(row=5+i+int(wp_mrocza), column=10).value = int(Lines_mr[i+l_art_mr*2])
    sheet.cell(row=5+i+int(wp_mrocza), column=6).value = Lines_mr[i+l_art_mr*3]
    if brak:
        sheet.cell(row=5+i+int(wp_mrocza), column=7).value = zdj[i]
    i = i + 1
sheet.cell(row=1, column=3).value = nr_gazety+"/2026 wersja_A"
source_file.save(fn)
#print("1x1 : ", str(sheet.cell(row=1, column=3).value))
source_file.close()
# SADKI
fn = ISSUE_OUTPUT_PREFIX + ""+nr_gazety+"/"+nr_gazety+"_Sd_IJ_01_rozpiska.xlsx"
source_file = openpyxl.load_workbook(fn)
sheet = source_file["Rozpiska"]
k=0
zdj = []
brak = False
for g in tab:
    if "Sd" in g:
        zdj.append(g[:-2])
        brak = True
    k = k + 1
i=0
while not i==l_art_sd:
    sheet.cell(row=5+i+int(wp_sadki), column=5).value = Lines_sd[i]
    sheet.cell(row=5+i+int(wp_sadki), column=3).value = Lines_sd[i+l_art_sd*1]
    sheet.cell(row=5+i+int(wp_sadki), column=10).value = int(Lines_sd[i+l_art_sd*2])
    sheet.cell(row=5+i+int(wp_sadki), column=6).value = Lines_sd[i+l_art_sd*3]
    if brak:
        sheet.cell(row=5+i+int(wp_sadki), column=7).value = zdj[i]
    i = i + 1
sheet.cell(row=1, column=3).value = nr_gazety+"/2026 wersja_A"
source_file.save(fn)
#print("1x1 : ", str(sheet.cell(row=1, column=3).value))
source_file.close()
# KCYNIA
fn = ISSUE_OUTPUT_PREFIX + ""+nr_gazety+"/"+nr_gazety+"_Kc_IJ_01_rozpiska.xlsx"
source_file = openpyxl.load_workbook(fn)
sheet = source_file["Rozpiska"]
k=0
zdj = []
brak = False
for g in tab:
    if "Kc" in g:
        zdj.append(g[:-2])
        brak = True
    k = k + 1
i=0
while not i==l_art_kc:
    sheet.cell(row=5+i+int(wp_kcynia), column=5).value = Lines_kc[i]
    sheet.cell(row=5+i+int(wp_kcynia), column=3).value = Lines_kc[i+l_art_kc*1]
    sheet.cell(row=5+i+int(wp_kcynia), column=10).value = int(Lines_kc[i+l_art_kc*2])
    sheet.cell(row=5+i+int(wp_kcynia), column=6).value = Lines_kc[i+l_art_kc*3]
    if brak:
        sheet.cell(row=5+i+int(wp_kcynia), column=7).value = str(zdj[i])
    i = i + 1
sheet.cell(row=1, column=3).value = nr_gazety+"/2026 wersja_A"
source_file.save(fn)
#print("1x1 : ", str(sheet.cell(row=1, column=3).value))
source_file.close()
# ROLNICZE
fn = ISSUE_OUTPUT_PREFIX + ""+nr_gazety+"/"+nr_gazety+"_RolA_IJ_01_rozpiska.xlsx"
source_file = openpyxl.load_workbook(fn)
sheet = source_file["Rozpiska"]
k=0
zdj = []
brak = False
for g in tab:
    if "RolA" in g:
        zdj.append(g[:-2])
        brak = True
    k = k + 1
i=0
while not i==l_art_rol:
    sheet.cell(row=5+i+int(wp_rol), column=5).value = Lines_rol[i]
    sheet.cell(row=5+i+int(wp_rol), column=3).value = Lines_rol[i+l_art_rol*1]
    sheet.cell(row=5+i+int(wp_rol), column=10).value = int(Lines_rol[i+l_art_rol*2])
    sheet.cell(row=5+i+int(wp_rol), column=6).value = Lines_rol[i+l_art_rol*3]
    if brak:
        sheet.cell(row=5+i+int(wp_rol), column=7).value = str(zdj[i])
    i = i + 1
sheet.cell(row=1, column=3).value = nr_gazety+"/2026 wersja_A"
source_file.save(fn)
#print("1x1 : ", str(sheet.cell(row=1, column=3).value))
source_file.close()
# POLICJA
fn = ISSUE_OUTPUT_PREFIX + ""+nr_gazety+"/"+nr_gazety+"_PolA_IJ_01_rozpiska.xlsx"
source_file = openpyxl.load_workbook(fn)
sheet = source_file["Rozpiska"]
k=0
zdj = []
brak = False
for g in tab:
    if "PolA" in g:
        zdj.append(g[:-2])
        brak = True
    k = k + 1
i=0
while not i==l_art_pol:
    sheet.cell(row=5+i+int(wp_pol), column=5).value = Lines_pol[i]
    sheet.cell(row=5+i+int(wp_pol), column=3).value = Lines_pol[i+l_art_pol*1]
    sheet.cell(row=5+i+int(wp_pol), column=10).value = int(Lines_pol[i+l_art_pol*2])
    sheet.cell(row=5+i+int(wp_pol), column=6).value = Lines_pol[i+l_art_pol*3]
    if brak:
        sheet.cell(row=5+i+int(wp_pol), column=7).value = str(zdj[i])
    i = i + 1
sheet.cell(row=1, column=3).value = nr_gazety+"/2026 wersja_A"
source_file.save(fn)
#print("1x1 : ", str(sheet.cell(row=1, column=3).value))
source_file.close()
